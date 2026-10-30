from __future__ import annotations

"""Serveur temps réel: MQTT (voiture) -> WebSocket (site web + app Flutter) + API REST.

Lancer (depuis le dossier server/):
    uvicorn server:app --host 0.0.0.0 --port 8000

Configuration modifiable dans: server/config.py
"""
import asyncio
import csv
import io
import json
import os
import sys
import time
import urllib.parse
import urllib.request
from collections import deque
from contextlib import asynccontextmanager
from pathlib import Path
from typing import Optional

import paho.mqtt.client as mqtt
from fastapi import FastAPI, HTTPException, Request, WebSocket, WebSocketDisconnect
from fastapi.responses import FileResponse, JSONResponse, Response
from fastapi.staticfiles import StaticFiles

import db
from alerts import AlertEngine
from fleet_sim import FleetSim

try:
    from config import (ALERT_ACK_TIMEOUT_S, EMERGENCY_MESSAGE, EMERGENCY_NUMBER, MAX_IMAGE_BYTES,
                        MQTT_HOST, MQTT_PORT, REPORT_INTERVAL_S, ROUTE_WAYPOINTS,
                        TELEGRAM_CHAT_ID, TELEGRAM_TOKEN)
except ImportError:  # sécurité: valeurs par défaut
    ALERT_ACK_TIMEOUT_S = 10.0
    EMERGENCY_NUMBER = "+216 00 000 000"
    EMERGENCY_MESSAGE = "Simulation: conducteur endormi, position envoyée à la Protection civile."
    MAX_IMAGE_BYTES = 2 * 1024 * 1024
    MQTT_HOST = os.getenv("MQTT_HOST", "localhost")
    MQTT_PORT = int(os.getenv("MQTT_PORT", "1883"))
    REPORT_INTERVAL_S = 120
    ROUTE_WAYPOINTS = [(35.8256, 10.5908), (35.8300, 10.6050), (35.8265, 10.6340), (35.8190, 10.6180)]
    TELEGRAM_TOKEN = os.getenv("TELEGRAM_TOKEN", "")
    TELEGRAM_CHAT_ID = os.getenv("TELEGRAM_CHAT_ID", "")

BASE = Path(__file__).resolve().parent
LOGS = BASE / "logs"
LOGS.mkdir(exist_ok=True)
STATIC = BASE / "static"
ROUTE_FILE = STATIC / "route.json"
FLEET_SIM = os.environ.get("FLEET_SIM", "1") != "0"  # 0 = ne pas faire rouler les voitures sans webcam

engine = AlertEngine()
# WebSocket clients: ws -> dict(user)
ws_clients: dict[WebSocket, dict] = {}
latest: dict[str, dict] = {}
recent_alerts: deque = deque(maxlen=30)
online: dict[str, bool] = {}
last_seen: dict[str, float] = {}
last_state_sent: dict[str, tuple[float, str]] = {}
mqtt_state = {"connected": False}
last_reports: dict[str, dict] = {}
alert_counts_since_report: dict[str, int] = {}
max_alert_level_since_report: dict[str, int] = {}
escalation_tasks: dict[int, asyncio.Task] = {}
_vcache: dict = {"t": 0.0, "by_code": {}}


# ---------- aides auth ----------
def get_token(request: Request) -> str:
    auth = request.headers.get("Authorization", "")
    if auth.startswith("Bearer "):
        return auth[7:].strip()
    if "token" in request.query_params:
        return request.query_params["token"]
    return ""


def get_current_user(request: Request) -> dict:
    token = get_token(request)
    if not token:
        raise HTTPException(status_code=401, detail="Non authentifié (token manquant)")
    user = db.get_user_by_token(token)
    if not user:
        raise HTTPException(status_code=401, detail="Session invalide ou expirée")
    return user


def require_supervisor(request: Request) -> dict:
    user = get_current_user(request)
    if user.get("role") != "supervisor":
        raise HTTPException(status_code=403, detail="Accès refusé: réservé au Superviseur (Police)")
    return user


def require_client(request: Request) -> dict:
    user = get_current_user(request)
    if user.get("role") != "client":
        raise HTTPException(status_code=403, detail="Accès réservé au conducteur (Client)")
    return user


# ---------- petites aides ----------
def log_csv(name: str, header: list, row: list) -> None:
    path = LOGS / name
    new = not path.exists()
    with open(path, "a", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        if new:
            w.writerow(header)
        w.writerow(row)


def vehicle_meta(code: str) -> dict:
    """Matricule + conducteur d'une voiture (petit cache de 10 s)."""
    now = time.time()
    if now - _vcache["t"] > 10:
        _vcache["by_code"] = {v["code"]: v for v in db.list_vehicles()}
        _vcache["t"] = now
    v = _vcache["by_code"].get(code)
    return {
        "plate": v["plate"] if v else code,
        "driver": (v or {}).get("driver_name") or "—",
        "cin": (v or {}).get("driver_cin") or "—",
        "phone": (v or {}).get("driver_phone") or "—",
        "emergency_name": (v or {}).get("emergency_name") or "—",
        "emergency_phone": (v or {}).get("emergency_phone") or "—",
    }


def invalidate_vehicle_cache() -> None:
    _vcache["t"] = 0.0


def runtime_settings() -> dict:
    s = db.get_settings()
    try:
        s["_ack"] = float(s.get("ack_timeout_s", ALERT_ACK_TIMEOUT_S))
    except ValueError:
        s["_ack"] = float(ALERT_ACK_TIMEOUT_S)
    return s


def apply_settings_to_engine() -> None:
    s = db.get_settings()
    try:
        engine.cfg.asleep_cooldown_s = float(s["asleep_cooldown_s"])
    except (KeyError, ValueError):
        pass


# ---------- Telegram (optionnel) ----------
def send_telegram(text: str) -> None:
    if not (TELEGRAM_TOKEN and TELEGRAM_CHAT_ID):
        return
    url = f"https://api.telegram.org/bot{TELEGRAM_TOKEN}/sendMessage"
    data = urllib.parse.urlencode({"chat_id": TELEGRAM_CHAT_ID, "text": text}).encode()
    try:
        urllib.request.urlopen(urllib.request.Request(url, data=data), timeout=10)
    except Exception as e:
        print("Erreur Telegram:", e)


def alert_text(a: dict) -> str:
    icons = {2: "⚠️", 3: "🚨", 4: "🆘"}
    lines = [f"{icons.get(a['level'], '⚠️')} {a['title']} (niveau {a['level']})",
             f"Véhicule: {a.get('plate') or a['vehicle_id']} - Conducteur: {a.get('driver', '—')}",
             a["message"]]
    if a.get("maps_url"):
        lines.append(f"Position: {a['maps_url']}")
    if a["level"] == 4:
        lines.append("SIMULATION: aucun appel réel n'est passé.")
    return "\n".join(lines)


# ---------- diffusion WebSocket filtrée par rôle ----------
async def broadcast(message: dict) -> None:
    text = json.dumps(message)
    dead = []
    target_vid = message.get("vehicle_id") or message.get("vehicle_code")
    for ws, user in list(ws_clients.items()):
        role = user.get("role")
        user_vid = user.get("vehicle_code")
        # Le superviseur reçoit TOUT
        # Le client reçoit uniquement les événements de SA voiture
        if role == "supervisor" or not user_vid or not target_vid or target_vid == user_vid:
            try:
                await ws.send_text(text)
            except Exception:
                dead.append(ws)
    for ws in dead:
        ws_clients.pop(ws, None)


async def send_report(vid: str) -> dict:
    now = time.time()
    v = latest.get(vid, {})
    lat, lon = v.get("lat"), v.get("lon")
    cnt = alert_counts_since_report.get(vid, 0)
    max_lvl = max_alert_level_since_report.get(vid, 0)
    alert_counts_since_report[vid] = 0
    max_alert_level_since_report[vid] = 0
    time_str = time.strftime("%H:%M:%S", time.localtime(now))
    maps_url = f"https://www.google.com/maps?q={lat},{lon}" if lat is not None else None
    meta = vehicle_meta(vid)
    rep = {
        "type": "report", "vehicle_id": vid, "plate": meta["plate"], "driver": meta["driver"],
        "ts": int(now * 1000), "time": time_str, "lat": lat, "lon": lon, "maps_url": maps_url,
        "speed": v.get("speed"), "state": v.get("state"),
        "alerts_count": cnt, "max_level": max_lvl,
    }
    last_reports[vid] = rep
    log_csv("reports.csv",
            ["time", "vehicle", "state", "speed", "lat", "lon", "alerts_count", "max_level"],
            [time_str, vid, rep["state"], rep["speed"], lat, lon, cnt, max_lvl])
    await broadcast(rep)
    if TELEGRAM_TOKEN and TELEGRAM_CHAT_ID:
        text = (f"📋 RAPPORT CONDUCTEUR ({time_str})\n"
                f"Véhicule: {meta['plate']} ({meta['driver']})\n"
                f"État: {rep['state']} | Vitesse: {rep['speed']} km/h\n"
                f"Alertes: {cnt} (max niveau {max_lvl})\n"
                f"Position: {maps_url or 'N/A'}")
        asyncio.get_running_loop().run_in_executor(None, send_telegram, text)
    return rep


def _fmt_pos(v) -> str:
    return f"{v:.4f}" if isinstance(v, (int, float)) else "?"


# ---------- escalade: personne ne répond -> appel police SIMULÉ ----------
async def escalate_later(det_id: int, vid: str, delay: float) -> None:
    try:
        await asyncio.sleep(delay)
        det = db.get_detection(det_id)
        if not det or det["status"] not in ("new", None):
            return
        db.set_detection_status(det_id, "escalated")
        with db.conn() as c:
            c.execute("UPDATE alerts SET status='police_called' WHERE detection_id=?", (det_id,))
        meta = vehicle_meta(vid)
        number = db.get_settings().get("emergency_number", EMERGENCY_NUMBER)
        msg = {
            "type": "escalation", "detection_id": det_id, "vehicle_id": vid,
            "plate": meta["plate"], "driver": meta["driver"], "cin": meta["cin"],
            "phone": meta["phone"], "number": number,
            "message": (f"Simulation: conducteur endormi, voiture {meta['plate']}, "
                        f"position {_fmt_pos(det.get('lat'))}, {_fmt_pos(det.get('lon'))}."),
            "ts": int(time.time() * 1000),
            "status": "police_called",
            "status_label": "Police appelée",
        }
        with db.conn() as c:
            arow = c.execute("SELECT id FROM alerts WHERE detection_id = ? ORDER BY id DESC LIMIT 1",
                             (det_id,)).fetchone()
        cur = latest.get(vid) or {}
        msg.update({
            "alert_id": arow["id"] if arow else None, "vehicle_code": vid,
            "image_url": det.get("image_url"), "lat": det.get("lat"), "lon": det.get("lon"),
            "eyes_closed_s": det.get("eyes_closed_s"), "speed": cur.get("speed"),
            "driver_name": meta["driver"], "driver_cin": meta["cin"], "driver_phone": meta["phone"],
            "emergency_name": meta["emergency_name"], "emergency_phone": meta["emergency_phone"],
        })
        print(f"ESCALADE (simulée): appel {number} pour {meta['plate']}")
        log_csv("alerts.csv", ["time", "vehicle", "level", "title", "message", "latency_ms", "simulated_call"],
                [round(time.time(), 3), vid, 4, "Escalade", "Police prévenue (SIMULÉ)", "", number])
        await broadcast(msg)
    except asyncio.CancelledError:
        pass
    finally:
        escalation_tasks.pop(det_id, None)


# ---------- traitement d'un état reçu (MQTT ou HTTP) ----------
async def handle_state(data: dict, force_broadcast: bool = False) -> None:
    vid = data.get("vehicle_id", "car01")
    now = time.time()
    latency = max(0, round(now * 1000 - data.get("ts", now * 1000)))
    meta = vehicle_meta(vid)
    data["type"] = "state"
    data["latency_ms"] = latency
    data["plate"] = meta["plate"]
    data["driver"] = meta["driver"]
    latest[vid] = data
    last_seen[vid] = now
    if not online.get(vid):
        online[vid] = True
        await broadcast({"type": "status", "vehicle_id": vid, "online": True})

    log_csv("states.csv",
            ["time", "vehicle", "state", "confidence", "latency_ms", "infer_ms", "lat", "lon"],
            [round(now, 3), vid, data.get("state"), data.get("confidence"), latency,
             data.get("infer_ms"), data.get("lat"), data.get("lon")])

    # Limite: 2 messages d'état par seconde et par voiture (sauf si l'état change)
    prev = last_state_sent.get(vid)
    if force_broadcast or prev is None or now - prev[0] >= 0.45 or prev[1] != data.get("state"):
        last_state_sent[vid] = (now, data.get("state"))
        await broadcast(data)

    loop = asyncio.get_running_loop()
    for a in engine.update(vid, data, now):
        a["plate"] = meta["plate"]
        a["driver"] = meta["driver"]
        a["cin"] = meta["cin"]
        a["phone"] = meta["phone"]
        a["image_url"] = None
        a["detection_id"] = None
        a["status"] = "new"
        a["status_label"] = "En attente"
        if a.get("kind") == "asleep":  # ligne dans l'historique des détections
            early = db.find_unlinked_detection(vid, a["ts"] - 10_000)
            if early:
                a["detection_id"] = early["id"]
                a["image_url"] = early["image_url"]
            else:
                a["detection_id"] = db.insert_detection(
                    vid, a["ts"], a.get("lat"), a.get("lon"), a.get("eyes_closed_s"), level=4)
            delay = runtime_settings()["_ack"]
            a["ack_timeout_s"] = delay
            escalation_tasks[a["detection_id"]] = asyncio.create_task(
                escalate_later(a["detection_id"], vid, delay))
        a["alert_id"] = db.insert_alert(a)
        recent_alerts.append(a)
        alert_counts_since_report[vid] = alert_counts_since_report.get(vid, 0) + 1
        max_alert_level_since_report[vid] = max(max_alert_level_since_report.get(vid, 0), a["level"])

        sim_call_str = json.dumps(a.get("simulated_call")) if a.get("simulated_call") else ""
        print(f"ALERTE niveau {a['level']}: {a['title']} - {a['message']} [{meta['plate']}]"
              + (f" | Appel simulé: {sim_call_str}" if sim_call_str else ""))
        log_csv("alerts.csv", ["time", "vehicle", "level", "title", "message", "latency_ms", "simulated_call"],
                [round(now, 3), vid, a["level"], a["title"], a["message"], latency, sim_call_str])
        await broadcast(a)

        if a["level"] >= 3:
            await send_report(vid)
        if a["notify"]:
            loop.run_in_executor(None, send_telegram, alert_text(a))


async def handle_status(vid: str, is_online: bool) -> None:
    online[vid] = is_online
    await broadcast({"type": "status", "vehicle_id": vid, "online": is_online})


# ---------- MQTT ----------
def start_mqtt(loop: asyncio.AbstractEventLoop) -> mqtt.Client:
    client = mqtt.Client(mqtt.CallbackAPIVersion.VERSION2, client_id="server-1")

    def on_connect(c, userdata, flags, reason_code, properties):
        mqtt_state["connected"] = True
        print(f"MQTT connecté à {MQTT_HOST}:{MQTT_PORT}")
        c.subscribe([("car/+/state", 0), ("car/+/status", 1)])

    def on_disconnect(c, userdata, flags, reason_code, properties):
        mqtt_state["connected"] = False
        print("MQTT déconnecté (reconnexion automatique)")

    def on_message(c, userdata, msg):
        parts = msg.topic.split("/")
        if len(parts) != 3:
            return
        try:
            if parts[2] == "state":
                data = json.loads(msg.payload)
                data.setdefault("vehicle_id", parts[1])
                asyncio.run_coroutine_threadsafe(handle_state(data), loop)
            elif parts[2] == "status":
                is_online = msg.payload.decode(errors="ignore") == "online"
                asyncio.run_coroutine_threadsafe(handle_status(parts[1], is_online), loop)
        except Exception as e:
            print("Message MQTT ignoré:", e)

    client.on_connect = on_connect
    client.on_disconnect = on_disconnect
    client.on_message = on_message
    try:
        client.connect_async(MQTT_HOST, MQTT_PORT, keepalive=30)
        client.loop_start()
    except Exception as e:
        print("MQTT indisponible (le serveur marche quand même en mode HTTP):", e)
    return client


@asynccontextmanager
async def lifespan(app: FastAPI):
    db.init_db(seed=True)
    apply_settings_to_engine()
    loop = asyncio.get_running_loop()
    client = start_mqtt(loop)

    async def report_worker():
        while True:
            await asyncio.sleep(REPORT_INTERVAL_S)
            for vid in list(latest.keys()):
                await send_report(vid)

    async def presence_worker():
        """Une voiture sans message depuis 8 s est considérée hors ligne."""
        while True:
            await asyncio.sleep(3)
            now = time.time()
            for vid, t in list(last_seen.items()):
                if online.get(vid) and now - t > 8:
                    online[vid] = False
                    await broadcast({"type": "status", "vehicle_id": vid, "online": False})

    async def fleet_worker():
        """Fait rouler les voitures SANS webcam qui ont un trajet dessiné par leur client (positions seulement)."""
        sim = FleetSim()
        while True:
            await asyncio.sleep(1.0)
            if not FLEET_SIM:
                continue
            try:
                now = time.time()
                for v in db.list_vehicles():
                    code = v.get("code")
                    if not code:
                        continue
                    cur = latest.get(code)
                    if cur and not cur.get("fleet_sim") and now - last_seen.get(code, 0) < 8:
                        continue  # une vraie webcam envoie des données: on ne simule pas
                    r = db.get_vehicle_route(code)
                    if not (r and r.get("points")):
                        continue  # seules les voitures dont le client a dessiné un trajet roulent
                    pts = r["points"]
                    speed = 38 + (sum(ord(ch) for ch in code) % 17)  # 38 à 54 km/h
                    pos = sim.step(code, pts, speed, 1.0)
                    if not pos:
                        continue
                    lat, lon, heading, progress = pos
                    await handle_state({
                        "vehicle_id": code, "state": "EyesOpen", "eyes_closed_s": 0.0,
                        "lat": round(lat, 6), "lon": round(lon, 6), "heading": round(heading, 1),
                        "speed": speed, "route_progress": round(progress, 1),
                        "ts": int(now * 1000), "fleet_sim": True,
                    })
            except Exception as e:  # ne jamais arrêter le serveur à cause de la simulation
                print("fleet_worker:", e)

    tasks = [asyncio.create_task(report_worker()), asyncio.create_task(presence_worker()),
             asyncio.create_task(fleet_worker())]
    yield
    for t in tasks:
        t.cancel()
    client.loop_stop()
    client.disconnect()


app = FastAPI(title="Driver Safety Monitor", lifespan=lifespan)

STATIC.mkdir(exist_ok=True)
(STATIC / "captures").mkdir(exist_ok=True)
app.mount("/static", StaticFiles(directory=STATIC), name="static")


def err(msg: str, code: int = 400) -> JSONResponse:
    return JSONResponse({"error": msg}, status_code=code)


@app.get("/")
async def index():
    return FileResponse(STATIC / "index.html", headers={"Cache-Control": "no-cache"})


@app.get("/brand-logo")
async def brand_logo():
    """Logo du site: static/brand.png|jpg|jpeg|webp|svg (celui du Bureau), sinon logo.svg."""
    for ext in ("png", "jpg", "jpeg", "webp", "svg"):
        f = STATIC / f"brand.{ext}"
        if f.exists():
            return FileResponse(f, headers={"Cache-Control": "no-cache"})
    return FileResponse(STATIC / "logo.svg", headers={"Cache-Control": "no-cache"})


# =================================================================== AUTHENTIFICATION
@app.post("/api/auth/register")
async def api_register(request: Request):
    """Inscription client (conducteur) avec génération automatique de code voiture et diffusion temps réel."""
    try:
        data = await request.json()
        plan_id = str(data.get("plan", "")).strip()
        if plan_id and not db.get_plan(plan_id):
            return err("Offre inconnue.", 400)
        res = db.register_client(data)
        if plan_id:
            res["subscription"] = db.create_subscription(
                res["user"].get("driver_id"), plan_id, str(data.get("card_last4", "")))
        invalidate_vehicle_cache()
        # Diffuser immédiatement l'événement WebSocket client_registered pour la police
        await broadcast({
            "type": "client_registered",
            "client": res["user"],
            "vehicle": {
                "id": res["user"].get("id"),
                "code": res["vehicle_code"],
                "plate": res["plate"],
                "brand": data.get("brand", ""),
                "model": data.get("model", ""),
                "color": data.get("color", "Blanche"),
                "driver_name": res["user"].get("full_name"),
                "driver_cin": res["user"].get("cin"),
                "driver_phone": res["user"].get("phone"),
                "driver_email": res["user"].get("email"),
                "emergency_name": res["user"].get("emergency_name"),
                "emergency_phone": res["user"].get("emergency_phone"),
                "online": False,
                "state": "Safe",
                "speed": 0,
                "alerts_count": 0
            },
            "route": res.get("route")
        })
        return res
    except ValueError as e:
        return err(str(e), 400)
    except Exception as e:
        return err(f"Erreur d'inscription: {e}", 500)


@app.post("/api/auth/login")
async def api_login(request: Request):
    """Connexion (Client ou Superviseur Police)."""
    try:
        data = await request.json()
        username = str(data.get("username", "")).strip()
        password = str(data.get("password", "")).strip()
        if not username or not password:
            return err("Identifiant et mot de passe requis", 400)
        user = db.authenticate_user(username, password)
        if not user:
            return err("Identifiant ou mot de passe incorrect", 401)
        token = db.create_session(user["id"])
        return {"ok": True, "token": token, "user": user}
    except Exception as e:
        return err(str(e), 500)


@app.get("/api/auth/me")
async def api_me(request: Request):
    """Récupère l'utilisateur connecté via l'en-tête Authorization: Bearer."""
    user = get_current_user(request)
    return {"ok": True, "user": user}


@app.post("/api/auth/logout")
async def api_logout(request: Request):
    """Déconnexion et invalidation du token."""
    token = get_token(request)
    if token:
        db.delete_session(token)
    return {"ok": True}


# =================================================================== ESPACE CLIENT
@app.get("/api/client/profile")
async def api_client_profile(request: Request):
    user = require_client(request)
    return user


@app.put("/api/client/profile")
async def api_client_profile_update(request: Request):
    user = require_client(request)
    data = await request.json()
    did = user.get("driver_id")
    if not did:
        return err("Conducteur introuvable", 404)
    db.update_client_profile(did, data)
    invalidate_vehicle_cache()
    fresh = db.get_user_by_id(user["id"])
    return fresh


# =================================================================== CAMÉRA DU TÉLÉPHONE
# Le téléphone envoie des images JPEG; le serveur fait la détection avec le MÊME moteur que
# detector_eyes.py (YOLO si models/eyes_best.pt existe, sinon MediaPipe + EAR).
import threading

DETECTOR_DIR = BASE.parent / "detector"
_phone = {"engine": None, "lock": threading.Lock(), "cars": {}}


def _phone_modules():
    if str(DETECTOR_DIR) not in sys.path:
        sys.path.insert(0, str(DETECTOR_DIR))
    import detector_eyes as de
    import utils as du
    return de, du


def _phone_engine():
    if _phone["engine"] is None:
        try:
            de, _ = _phone_modules()
            model = BASE.parent / "models" / "eyes_best.pt"
            _phone["engine"] = de.YoloEyes(str(model), 0.4, 640) if model.exists() else de.MediaPipeEyes(0.21)
        except (SystemExit, ImportError) as e:
            raise RuntimeError("Moteur de détection indisponible sur le serveur. Installe: "
                               "pip install opencv-python mediapipe==0.10.14 (" + str(e) + ")")
    return _phone["engine"]


def _process_phone_frame(vid: str, jpeg: bytes, limit_s: float, cooldown_s: float, route_pts) -> tuple:
    import cv2
    import numpy as np
    de, du = _phone_modules()
    frame = cv2.imdecode(np.frombuffer(jpeg, np.uint8), cv2.IMREAD_COLOR)
    if frame is None:
        raise ValueError("image illisible")
    car = _phone["cars"].setdefault(vid, {"timer": None, "route": None, "last": time.time(),
                                           "fps": 0.0, "last_capture": -1e9,
                                           "best": None, "captured": False})
    with _phone["lock"]:
        eng = _phone_engine()
        t0 = time.time()
        dets = eng.detect(frame)
        infer_ms = (time.time() - t0) * 1000
    now = time.time()
    n_open = sum(1 for d in dets if d[0] == "open")
    n_closed = sum(1 for d in dets if d[0] == "closed")
    if car["timer"] is None or abs(car["timer"].closed_seconds - limit_s) > 1e-6:
        car["timer"] = du.EyeTimer(closed_seconds=limit_s)
    if car["route"] is None:
        pts = [[float(p[0]), float(p[1])] for p in (route_pts or [])]
        car["route"] = du.RouteSimulator(points=pts if len(pts) >= 2 else None, speed_kmh=du.DEFAULT_SPEED_KMH)
    state, closed_s = car["timer"].update(n_open, n_closed, now)
    dt = max(now - car["last"], 1e-3)
    car["last"] = now
    car["fps"] = 0.8 * car["fps"] + 0.2 * (1.0 / dt)
    lat, lon, heading, progress = car["route"].step(min(dt, 2.0))
    confs = [d[1] for d in dets]
    payload = {
        "vehicle_id": vid, "state": state, "eyes_closed_s": round(closed_s, 2),
        "eyes_open": n_open, "eyes_closed": n_closed,
        "confidence": round(sum(confs) / len(confs), 3) if confs else 0.0,
        "lat": round(lat, 6), "lon": round(lon, 6), "heading": round(heading, 1),
        "route_progress": round(progress, 1), "speed": du.DEFAULT_SPEED_KMH, "impact": False,
        "ts": int(now * 1000), "infer_ms": round(infer_ms, 1), "fps": round(car["fps"], 1),
        "source": "phone",
    }
    # Meilleure image de l'épisode (yeux fermés avec la meilleure confiance), pour l'envoyer avec l'alerte
    closed_now = n_closed >= 1 and n_open == 0
    if closed_now:
        score = sum(d[1] for d in dets if d[0] == "closed") / n_closed
        if car["best"] is None or score > car["best"][0]:
            car["best"] = (score, frame.copy(), list(dets))
    # nouvel épisode seulement après 1 s d'yeux ouverts (comme le moteur d'alertes)
    if n_open > 0:
        if car.get("open_since") is None:
            car["open_since"] = now
        if now - car["open_since"] >= 1.0:
            car["best"], car["captured"] = None, False
    else:
        car["open_since"] = None
    capture = None
    if state == "Asleep" and (not car["captured"] or now - car["last_capture"] >= cooldown_s):
        car["captured"], car["last_capture"] = True, now
        src_frame, src_dets = (car["best"][1], car["best"][2]) if car["best"] else (frame, dets)
        de.draw_overlay(src_frame, src_dets, state, closed_s, limit_s, car["fps"], lat, lon, False)
        ok, buf = cv2.imencode(".jpg", src_frame, [int(cv2.IMWRITE_JPEG_QUALITY), 80])
        if ok:
            capture = buf.tobytes()
    h, w = frame.shape[:2]
    return payload, dets, capture, (w, h)


async def _store_phone_capture(vid: str, jpeg: bytes, payload: dict) -> None:
    """Enregistre l'image AVANT la création de l'alerte: l'alerte la récupère (find_unlinked_detection)."""
    now = time.time()
    day = time.strftime("%Y-%m-%d", time.localtime(now))
    folder = db.CAPTURES_DIR / day
    folder.mkdir(parents=True, exist_ok=True)
    name = f"{vid}_{time.strftime('%H%M%S', time.localtime(now))}_{int(now * 1000) % 1000:03d}.jpg"
    (folder / name).write_bytes(jpeg)
    rel = f"captures/{day}/{name}"
    det_id = db.insert_detection(vid, int(now * 1000), payload.get("lat"), payload.get("lon"),
                                 payload.get("eyes_closed_s"), level=4, image_path=rel)
    d = db.get_detection(det_id)
    await broadcast({"type": "detection_update", "detection": d, "detection_id": det_id,
                     "vehicle_id": vid, "image_url": d["image_url"]})


@app.post("/api/client/frame")
async def api_client_frame(request: Request):
    """Reçoit une image JPEG de la caméra du téléphone du conducteur connecté et renvoie le résultat."""
    user = require_client(request)
    vid = user.get("vehicle_code")
    body = await request.body()
    if len(body) > MAX_IMAGE_BYTES:
        return err("image trop grande", 413)
    if len(body) < 4 or body[:3] != b"\xff\xd8\xff":
        return err("seulement du JPEG", 415)
    st = db.get_settings()
    try:
        limit_s = float(st.get("eyes_closed_s", 5))
    except (TypeError, ValueError):
        limit_s = 5.0
    try:
        cooldown_s = float(st.get("asleep_cooldown_s", 30))
    except (TypeError, ValueError):
        cooldown_s = 30.0
    route = db.get_vehicle_route(vid)
    loop = asyncio.get_running_loop()
    try:
        payload, dets, capture, (w, h) = await loop.run_in_executor(
            None, _process_phone_frame, vid, body, limit_s, cooldown_s, (route or {}).get("points"))
    except RuntimeError as e:
        return err(str(e), 503)
    except Exception as e:
        return err(f"Analyse impossible: {e}", 500)
    if capture:   # image enregistrée d'abord: elle part avec l'alerte (plus d'image noire)
        await _store_phone_capture(vid, capture, payload)
    await handle_state(payload)
    return {"ok": True, "state": payload["state"], "eyes_closed_s": payload["eyes_closed_s"],
            "eyes_open": payload["eyes_open"], "eyes_closed": payload["eyes_closed"],
            "infer_ms": payload["infer_ms"], "w": w, "h": h,
            "boxes": [[d[0], round(d[1], 2)] + [round(v, 1) for v in d[2]] for d in dets]}


@app.get("/api/plans")
async def api_plans():
    """Offres proposées (affichage seulement)."""
    return db.PLANS


@app.get("/api/client/subscription")
async def api_client_subscription(request: Request):
    user = require_client(request)
    return db.get_subscription(user.get("driver_id")) or {}


@app.post("/api/client/subscription")
async def api_client_subscription_set(request: Request):
    """Souscription / changement d'offre. Paiement SIMULÉ: on ne reçoit que les 4 derniers chiffres."""
    user = require_client(request)
    data = await request.json()
    try:
        return db.create_subscription(user.get("driver_id"), str(data.get("plan", "")),
                                      str(data.get("card_last4", "")))
    except ValueError as e:
        return err(str(e), 400)


@app.get("/api/client/state")
async def api_client_state(request: Request):
    user = require_client(request)
    vid = user.get("vehicle_code")
    st = latest.get(vid, {"vehicle_id": vid, "state": "Safe", "speed": 0})
    return {
        "vehicle_code": vid,
        "plate": user.get("plate"),
        "online": bool(online.get(vid)),
        "state": st.get("state", "Safe"),
        "speed": st.get("speed", 0),
        "lat": st.get("lat"),
        "lon": st.get("lon"),
        "eyes_closed_s": st.get("eyes_closed_s", 0),
        "last_seen": last_seen.get(vid),
    }


@app.post("/api/client/route")
async def api_client_route_save(request: Request):
    """Enregistre le tracé interactif dessiné par le conducteur."""
    user = require_client(request)
    vid = user.get("vehicle_code")
    data = await request.json()
    wps = data.get("waypoints", [])
    pts = data.get("points", [])
    dist = float(data.get("distance_km", 0) or 0)
    if not wps or not pts or len(pts) < 2:
        return err("Tracé invalide: au moins 2 points requis", 400)
    saved = db.save_vehicle_route(vid, wps, pts, dist)
    await broadcast({
        "type": "route_updated",
        "vehicle_code": vid,
        "plate": user.get("plate"),
        "route": saved
    })
    return {"ok": True, "route": saved}


@app.get("/api/client/route")
async def api_client_route(request: Request):
    """Renvoie le tracé personnalisé du conducteur ou le tracé par défaut."""
    user = require_client(request)
    vid = user.get("vehicle_code")
    saved = db.get_vehicle_route(vid)
    if saved:
        cur = latest.get(vid)
        saved["current_pos"] = [cur["lat"], cur["lon"]] if cur and cur.get("lat") else None
        saved["plate"] = user.get("plate")
        return saved
    try:
        route_data = json.loads(ROUTE_FILE.read_text(encoding="utf-8"))
    except Exception:
        route_data = {"fallback": True, "waypoints": [list(p) for p in ROUTE_WAYPOINTS],
                      "points": [list(p) for p in ROUTE_WAYPOINTS]}
    cur = latest.get(vid)
    route_data["current_pos"] = [cur["lat"], cur["lon"]] if cur and cur.get("lat") else None
    route_data["vehicle_code"] = vid
    route_data["plate"] = user.get("plate")
    return route_data


@app.get("/api/vehicles/{code}/route")
async def api_vehicle_route_public(code: str):
    """Endpoint PUBLIC (sans login) pour detector_eyes.py --vehicle {code}."""
    saved = db.get_vehicle_route(code)
    if saved:
        return saved
    try:
        route_data = json.loads(ROUTE_FILE.read_text(encoding="utf-8"))
        route_data["vehicle_code"] = code
        return route_data
    except Exception:
        return {"vehicle_code": code, "fallback": True,
                "waypoints": [list(p) for p in ROUTE_WAYPOINTS],
                "points": [list(p) for p in ROUTE_WAYPOINTS]}


@app.post("/api/route/snap")
async def api_route_snap(request: Request):
    """Relie les points par de vraies routes (OSRM) avec secours interpolé en lignes droites."""
    data = await request.json()
    waypoints = data.get("waypoints", [])
    if not waypoints or len(waypoints) < 2:
        return err("Au moins 2 points requis", 400)
    # 1. Tenter OSRM public
    try:
        coords = ";".join(f"{float(p[1])},{float(p[0])}" for p in waypoints)
        url = f"https://router.project-osrm.org/route/v1/driving/{coords}?overview=full&geometries=geojson&continue_straight=true"
        req = urllib.request.Request(url, headers={"User-Agent": "driver-safety-monitor/1.0"})
        loop = asyncio.get_running_loop()
        def _fetch():
            with urllib.request.urlopen(req, timeout=7) as r:
                return json.loads(r.read().decode())
        res = await loop.run_in_executor(None, _fetch)
        if res.get("code") == "Ok" and res.get("routes"):
            route_obj = res["routes"][0]
            points = [[round(lat, 6), round(lon, 6)] for lon, lat in route_obj["geometry"]["coordinates"]]
            dist_km = round(route_obj["distance"] / 1000.0, 2)
            return {"ok": True, "points": points, "distance_km": dist_km, "osrm": True}
    except Exception as e:
        print("OSRM non accessible, calcul en lignes droites:", e)

    # 2. Secours: interpolation en ligne droite (25 mètres par pas)
    import math
    def _d(a, b):
        r = 6371000.0
        la1, lo1, la2, lo2 = map(math.radians, (a[0], a[1], b[0], b[1]))
        dla, dlo = la2 - la1, lo2 - lo1
        h = math.sin(dla / 2) ** 2 + math.cos(la1) * math.cos(la2) * math.sin(dlo / 2) ** 2
        return 2 * r * math.asin(math.sqrt(h))

    pts = []
    total_m = 0.0
    for a, b in zip(waypoints[:-1], waypoints[1:]):
        pa = (float(a[0]), float(a[1]))
        pb = (float(b[0]), float(b[1]))
        dm = _d(pa, pb)
        total_m += dm
        n = max(1, int(dm / 25.0))
        for i in range(n):
            f = i / n
            pts.append([round(pa[0] + (pb[0] - pa[0]) * f, 6), round(pa[1] + (pb[1] - pa[1]) * f, 6)])
    pts.append([round(float(waypoints[-1][0]), 6), round(float(waypoints[-1][1]), 6)])
    return {"ok": True, "points": pts, "distance_km": round(total_m / 1000.0, 2), "osrm": False}


@app.get("/api/client/alerts")
async def api_client_alerts(request: Request, limit: int = 100):
    user = require_client(request)
    return db.list_alerts(vehicle_code=user.get("vehicle_code"), limit=limit)


@app.get("/api/client/detections")
async def api_client_detections(request: Request, limit: int = 500):
    user = require_client(request)
    return db.list_detections(vehicle=user.get("vehicle_code"), limit=limit)


@app.post("/api/client/alerts/{alert_id}/awake")
async def api_client_awake(alert_id: int, request: Request):
    user = require_client(request)
    al = db.get_alert(alert_id)
    if not al or al.get("vehicle_code") != user.get("vehicle_code"):
        return err("Alerte introuvable", 404)
    updated = db.update_alert_status(alert_id, "cancelled")
    det_id = al.get("detection_id")
    if det_id:
        task = escalation_tasks.pop(det_id, None)
        if task:
            task.cancel()
        db.set_detection_status(det_id, "acknowledged")
    await broadcast({
        "type": "ack", "detection_id": det_id, "vehicle_id": user.get("vehicle_code"),
        "alert_id": alert_id, "status": "cancelled", "status_label": "Annulé par moi"
    })
    return {"ok": True, "alert": updated}


# =================================================================== ESPACE SUPERVISEUR (POLICE)
@app.get("/api/status")
async def status():
    return {"mqtt_connected": mqtt_state["connected"], "vehicles": latest,
            "online": online, "alerts": list(recent_alerts),
            "reports": list(last_reports.values())}


@app.get("/api/route")
async def route():
    try:
        return JSONResponse(json.loads(ROUTE_FILE.read_text(encoding="utf-8")))
    except Exception:
        return {"fallback": True, "waypoints": [list(p) for p in ROUTE_WAYPOINTS],
                "points": [list(p) for p in ROUTE_WAYPOINTS]}


@app.get("/api/settings")
async def get_settings():
    s = db.get_settings()
    s["server_time"] = int(time.time() * 1000)
    return s


@app.put("/api/settings")
async def put_settings(request: Request):
    require_supervisor(request)
    body = await request.json()
    clean = {}
    for k in ("eyes_closed_s", "asleep_cooldown_s", "ack_timeout_s"):
        if k in body:
            try:
                v = float(body[k])
            except (TypeError, ValueError):
                return err(f"{k}: nombre attendu")
            if not (0.5 <= v <= 3600):
                return err(f"{k}: entre 0.5 et 3600")
            clean[k] = v
    if "emergency_number" in body:
        clean["emergency_number"] = str(body["emergency_number"])[:32]
    s = db.set_settings(clean)
    apply_settings_to_engine()
    return s


# ---------- voitures ----------
@app.get("/api/vehicles")
async def api_vehicles():
    out = db.list_vehicles()
    for v in out:
        v["online"] = bool(online.get(v["code"]))
        s = latest.get(v["code"])
        v["state"] = s.get("state") if s else None
        v["lat"] = s.get("lat") if s else None
        v["lon"] = s.get("lon") if s else None
        v["speed"] = s.get("speed") if s else 0
    return out


@app.post("/api/vehicles")
async def api_vehicle_add(request: Request):
    require_supervisor(request)
    try:
        v = db.save_vehicle(await request.json())
    except Exception as e:
        return err(str(e))
    invalidate_vehicle_cache()
    return v


@app.put("/api/vehicles/{vid}")
async def api_vehicle_edit(vid: int, request: Request):
    require_supervisor(request)
    try:
        v = db.save_vehicle(await request.json(), vid)
    except Exception as e:
        return err(str(e))
    invalidate_vehicle_cache()
    return v


@app.delete("/api/vehicles/{vid}")
async def api_vehicle_del(vid: int, request: Request):
    require_supervisor(request)
    db.delete_vehicle(vid)
    invalidate_vehicle_cache()
    return {"ok": True}


# ---------- conducteurs / clients ----------
@app.get("/api/drivers")
async def api_drivers():
    return db.list_drivers()


@app.post("/api/drivers")
async def api_driver_add(request: Request):
    require_supervisor(request)
    try:
        return db.save_driver(await request.json())
    except Exception as e:
        return err(str(e))


@app.put("/api/drivers/{did}")
async def api_driver_edit(did: int, request: Request):
    require_supervisor(request)
    try:
        d = db.save_driver(await request.json(), did)
    except Exception as e:
        return err(str(e))
    invalidate_vehicle_cache()
    return d


@app.delete("/api/drivers/{did}")
async def api_driver_del(did: int, request: Request):
    require_supervisor(request)
    db.delete_driver(did)
    invalidate_vehicle_cache()
    return {"ok": True}


# ---------- détections (images) ----------
@app.get("/api/detections")
async def api_detections(vehicle: str = "", status: str = "", date_from: str = "",
                         date_to: str = "", limit: int = 500):
    try:
        return db.list_detections(vehicle or None, date_from or None, date_to or None,
                                  status or None, min(max(limit, 1), 2000))
    except ValueError:
        return err("date au format AAAA-MM-JJ")


@app.get("/api/detections/{det_id}")
async def api_detection(det_id: int):
    d = db.get_detection(det_id)
    return d if d else err("introuvable", 404)


# NOTE IMPORTANTE: Cet endpoint RESTE SANS LOGIN pour que la webcam fonctionne directement
@app.post("/api/detections")
async def api_detection_upload(request: Request, vehicle: str = "car01", lat: Optional[float] = None,
                               lon: Optional[float] = None, eyes_closed_s: Optional[float] = None,
                               ts: Optional[int] = None):
    """Le détecteur webcam envoie ici l'image JPEG (corps de la requête = les octets du JPEG)."""
    body = await request.body()
    if len(body) > MAX_IMAGE_BYTES:
        return err("image trop grande", 413)
    if len(body) < 4 or body[:3] != b"\xff\xd8\xff":
        return err("seulement du JPEG", 415)
    vehicle = "".join(ch for ch in vehicle if ch.isalnum() or ch in "-_")[:32] or "car01"
    now = time.time()
    ts_ms = int(ts) if ts else int(now * 1000)
    day = time.strftime("%Y-%m-%d", time.localtime(now))
    folder = db.CAPTURES_DIR / day
    folder.mkdir(parents=True, exist_ok=True)
    name = f"{vehicle}_{time.strftime('%H%M%S', time.localtime(now))}_{int(now * 1000) % 1000:03d}.jpg"
    (folder / name).write_bytes(body)
    rel = f"captures/{day}/{name}"

    det = db.find_pending_detection(vehicle, int(now * 1000) - 60_000)
    if det:
        db.set_detection_image(det["id"], rel)
        det_id = det["id"]
    else:  # l'image arrive avant l'alerte: on crée la ligne
        det_id = db.insert_detection(vehicle, ts_ms, lat, lon, eyes_closed_s, level=4, image_path=rel)
    d = db.get_detection(det_id)
    await broadcast({"type": "detection_update", "detection": d, "detection_id": det_id,
                     "vehicle_id": vehicle, "image_url": d["image_url"]})
    return {"ok": True, "id": det_id, "image_url": d["image_url"]}


@app.post("/api/detections/{det_id}/ack")
async def api_detection_ack(det_id: int):
    d = db.get_detection(det_id)
    if not d:
        return err("introuvable", 404)
    db.set_detection_status(det_id, "acknowledged")
    with db.conn() as c:
        c.execute("UPDATE alerts SET status='cancelled' WHERE detection_id=?", (det_id,))
    task = escalation_tasks.pop(det_id, None)
    if task:
        task.cancel()
    await broadcast({"type": "ack", "detection_id": det_id, "vehicle_id": d["vehicle_code"],
                     "status": "cancelled", "status_label": "Annulé par moi"})
    return {"ok": True}


# ---------- suppression des détections (images) ----------
def _det_matches_scope(d: dict, scope: str) -> bool:
    """scope: all | missing (image introuvable) | demo (images de démonstration)."""
    path = d.get("image_path")
    if scope == "all":
        return True
    if scope == "demo":
        return bool(path) and "_demo" in str(path)
    if scope == "missing":
        if not path:
            return True
        return not (STATIC / path).exists()
    return False


async def _delete_detections(vehicle: Optional[str], scope: str) -> int:
    n = 0
    for d in db.list_detections(vehicle or None, None, None, None, 5000):
        if not _det_matches_scope(d, scope):
            continue
        task = escalation_tasks.pop(d["id"], None)
        if task:
            task.cancel()
        db.delete_detection(d["id"])
        n += 1
    return n


@app.delete("/api/detections/{det_id}")
async def api_detection_del(det_id: int, request: Request):
    require_supervisor(request)
    task = escalation_tasks.pop(det_id, None)
    if task:
        task.cancel()
    db.delete_detection(det_id)
    return {"ok": True}


@app.delete("/api/detections")
async def api_detections_del_many(request: Request, scope: str = "all", vehicle: str = ""):
    """Police: supprime plusieurs détections (scope = all | missing | demo)."""
    require_supervisor(request)
    if scope not in ("all", "missing", "demo"):
        return err("scope invalide (all, missing, demo)")
    n = await _delete_detections(vehicle or None, scope)
    await broadcast({"type": "detections_deleted", "vehicle_id": vehicle or None})
    return {"ok": True, "deleted": n}


@app.delete("/api/client/detections/{det_id}")
async def api_client_detection_del(det_id: int, request: Request):
    """Client: supprime UNE de ses détections (image + ligne)."""
    user = require_client(request)
    d = db.get_detection(det_id)
    if not d or d.get("vehicle_code") != user.get("vehicle_code"):
        return err("Détection introuvable", 404)
    task = escalation_tasks.pop(det_id, None)
    if task:
        task.cancel()
    db.delete_detection(det_id)
    return {"ok": True, "deleted": 1}


@app.delete("/api/client/detections")
async def api_client_detections_del_many(request: Request, scope: str = "all"):
    """Client: supprime plusieurs de SES détections (scope = all | missing | demo)."""
    user = require_client(request)
    if scope not in ("all", "missing", "demo"):
        return err("scope invalide (all, missing, demo)")
    n = await _delete_detections(user.get("vehicle_code"), scope)
    return {"ok": True, "deleted": n}


# ---------- alertes et statistiques ----------
@app.get("/api/alerts")
async def api_alerts(limit: int = 100):
    return db.list_alerts(limit=min(max(limit, 1), 500))


@app.post("/api/alerts/{alert_id}/action")
async def api_alert_action(alert_id: int, request: Request):
    """Actions police: 'respond' (a répondu), 'dispatch' (envoyer unité), 'close' (clôturer)."""
    require_supervisor(request)
    data = await request.json()
    action = data.get("action")
    status_map = {
        "respond": "responded",
        "dispatch": "dispatched",
        "close": "closed",
    }
    if action not in status_map:
        return err("Action invalide ('respond', 'dispatch', 'close')", 400)
    new_st = status_map[action]
    updated = db.update_alert_status(alert_id, new_st)
    if not updated:
        return err("Alerte introuvable", 404)
    vid = updated.get("vehicle_code")
    await broadcast({
        "type": "alert_status",
        "alert_id": alert_id,
        "detection_id": updated.get("detection_id"),
        "vehicle_id": vid,
        "status": new_st,
        "status_label": updated.get("status_label"),
        "action": action
    })
    return {"ok": True, "alert": updated}


@app.get("/api/stats")
async def api_stats():
    s = db.stats()
    vehicles = db.list_vehicles()
    s["vehicles_online"] = sum(1 for v in vehicles if online.get(v["code"]))
    s["vehicles_offline"] = s["vehicles_total"] - s["vehicles_online"]
    return s


# ---------- export Excel ----------
def _xlsx(title: str, header: list[str], data: list[list]) -> Response:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = title[:30]
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0F9D6B")
        c.alignment = Alignment(horizontal="center")
    for row in data:
        ws.append(row)
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 60)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{title}.xlsx"'})


@app.get("/api/export/vehicles.xlsx")
async def export_vehicles(request: Request):
    require_supervisor(request)
    data = [[v["plate"], v["driver_name"] or "—", v["brand"], v["model"], v["color"],
             v["code"], v["alerts_count"]] for v in db.list_vehicles()]
    return _xlsx("voitures", ["Matricule", "Nom du conducteur", "Marque", "Modèle", "Couleur",
                              "Code voiture", "Nombre d'alertes"], data)


@app.get("/api/export/detections.xlsx")
async def export_detections(request: Request, vehicle: str = "", status: str = "",
                            date_from: str = "", date_to: str = ""):
    require_supervisor(request)
    try:
        dets = db.list_detections(vehicle or None, date_from or None, date_to or None,
                                  status or None, 5000)
    except ValueError:
        return err("date au format AAAA-MM-JJ")
    base = str(request.base_url).rstrip("/")
    data = []
    for d in dets:
        when = time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(d["ts"] / 1000))
        pos = f"{d['lat']:.5f}, {d['lon']:.5f}" if d.get("lat") is not None else ""
        data.append([when, d.get("plate") or d["vehicle_code"], d.get("driver_name") or "—", pos,
                     d.get("eyes_closed_s"), d["status"],
                     f"{base}{d['image_url']}" if d.get("image_url") else ""])
    return _xlsx("detections", ["Date", "Matricule", "Nom du conducteur", "Position",
                                "Yeux fermés (s)", "Statut", "Image"], data)


# ---------- entrées webcam & tests (SANS LOGIN POUR DETECTEUR) ----------
@app.post("/api/state")
async def api_state(request: Request):
    """Même chose que MQTT mais en HTTP (utilisé par detector_eyes.py --transport http)."""
    data = await request.json()
    if not isinstance(data, dict) or "state" not in data:
        return err("champ 'state' obligatoire")
    data.setdefault("vehicle_id", "car01")
    await handle_state(data)
    return {"ok": True}


@app.post("/api/simulate/asleep")
async def api_simulate_asleep(vehicle: str = "car01"):
    """Bouton 'Tester l'alerte': simule une alerte endormi avec image pour tester le flux."""
    try:
        from demo_image import make_demo_jpeg
    except Exception as e:
        return err(f"opencv manquant pour l'image de démo: {e}", 500)
    base = latest.get(vehicle) or {"vehicle_id": vehicle, "lat": 35.8256, "lon": 10.5908, "speed": 45}
    msg = dict(base)
    msg.update({"vehicle_id": vehicle, "state": "Asleep", "eyes_closed_s": 5.0,
                "ts": int(time.time() * 1000), "simulated": True})
    engine.reset_asleep_cooldown(vehicle)
    await handle_state(msg, force_broadcast=True)
    jpeg = make_demo_jpeg(5.0, vehicle)
    now = time.time()
    day = time.strftime("%Y-%m-%d", time.localtime(now))
    folder = db.CAPTURES_DIR / day
    folder.mkdir(parents=True, exist_ok=True)
    name = f"{vehicle}_{time.strftime('%H%M%S', time.localtime(now))}_demo.jpg"
    (folder / name).write_bytes(jpeg)
    rel = f"captures/{day}/{name}"
    det = db.find_pending_detection(vehicle, int(now * 1000) - 60_000)
    if det:
        db.set_detection_image(det["id"], rel)
        d = db.get_detection(det["id"])
        await broadcast({"type": "detection_update", "detection": d, "detection_id": d["id"],
                         "vehicle_id": vehicle, "image_url": d["image_url"]})
        return {"ok": True, "detection_id": d["id"]}
    return {"ok": True}


# ---------- WebSocket avec authentification et filtrage ----------
@app.websocket("/ws")
async def ws_endpoint(ws: WebSocket):
    token = ws.query_params.get("token")
    user = db.get_user_by_token(token) if token else None
    await ws.accept()

    # Si aucun token fourni (ex: navigateur non connecté), invité superviseur pour vue par défaut
    if not user:
        user = {"role": "supervisor", "username": "guest"}

    ws_clients[ws] = user
    try:
        # Snapshot personnalisé selon le rôle
        if user.get("role") == "client":
            my_vid = user.get("vehicle_code")
            st = [latest[my_vid]] if my_vid in latest else []
            al = [a for a in recent_alerts if a.get("vehicle_id") == my_vid or a.get("vehicle_code") == my_vid]
            rep = [r for r in last_reports.values() if r.get("vehicle_id") == my_vid]
            on = {my_vid: online.get(my_vid, False)} if my_vid else {}
            await ws.send_text(json.dumps({
                "type": "snapshot",
                "role": "client",
                "states": st,
                "alerts": al,
                "reports": rep,
                "online": on,
                "my_vehicle": my_vid,
                "plate": user.get("plate"),
                "user": user,
            }))
        else:
            await ws.send_text(json.dumps({
                "type": "snapshot",
                "role": "supervisor",
                "states": list(latest.values()),
                "alerts": list(recent_alerts),
                "reports": list(last_reports.values()),
                "online": online,
                "user": user,
            }))

        while True:
            await ws.receive_text()
    except Exception:
        pass
    finally:
        ws_clients.pop(ws, None)
